"""Export academic validation report from generated synthetic data to an Excel (.xlsx) workbook.

This script produces a clean, professional synthetic data report for 3 educational programs:
- Sheet 1: 'Đại trà' (Standard Program)
- Sheet 2: 'Chất lượng cao' (High-Quality Program)
- Sheet 3: 'Tiên tiến' (Advanced Program)

Structure for each sheet:
1. Tiêu đề & Bối cảnh: Báo cáo tổng quan dữ liệu mô phỏng phục vụ mô hình Cảnh báo học vụ sớm (EWS).
2. BẢNG 1: Tỷ lệ Tốt nghiệp đúng hạn (Khóa K20, K21, K22).
   - Dòng hướng dẫn công thức / logic tính metric trên cơ sở dữ liệu thực tế.
3. BẢNG 2: Tỷ lệ Điểm & Xếp loại của những sinh viên đã tốt nghiệp (Khóa K20, K21, K22).
   - GPA TB Hệ 10, CPA TB Hệ 4 và tỷ lệ % phân bố 6 mức xếp loại điểm.
   - Dòng hướng dẫn công thức / logic tính trên dữ liệu thực tế.
4. BIỂU ĐỒ TRỰC QUAN CỦA BẢNG 2 (Khung A19):
   - Biểu đồ Cột chồng 100% (100% Stacked Column Chart) so sánh cơ cấu điểm tốt nghiệp K20, K21, K22.
5. BẢNG 3: Tỷ lệ Bỏ học từng khóa (Khóa K20 đến K26 - Khung A44).
   - Sĩ số tuyển sinh, số lượng sinh viên bỏ học và tỷ lệ % bỏ học.
   - Dòng hướng dẫn công thức / logic tính trên dữ liệu thực tế.
"""

from pathlib import Path
from typing import Any, Dict, List, Tuple
import numpy as np
import openpyxl
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

ROOT_DIR = Path(__file__).resolve().parents[1]
METRICS_DIR = ROOT_DIR / "generated" / "metrics"
USERS_DIR = ROOT_DIR / "generated" / "users"
DEFAULT_OUTPUT_FILE = ROOT_DIR / "academic_validation_report.xlsx"


def load_all_datasets() -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Load student metrics and users datasets, computing cumulative GPA 10 from enrollments."""
    student_file = METRICS_DIR / "student_metrics.csv"
    user_student_file = USERS_DIR / "student.csv"
    courses_file = ROOT_DIR / "generated" / "academics" / "courses.csv"
    enrollment_dir = ROOT_DIR / "generated" / "enrollments"

    if not student_file.exists():
        raise FileNotFoundError(f"Required file not found: {student_file}")

    df_students = pd.read_csv(student_file, dtype={"student_id": str, "cohort": str})
    df_user_st = pd.read_csv(user_student_file, dtype={"user_id": str}) if user_student_file.exists() else pd.DataFrame()

    if not df_user_st.empty and "program" in df_user_st.columns:
        df_students = df_students.merge(df_user_st[["user_id", "program"]], left_on="student_id", right_on="user_id", how="left")

    # Use precomputed gpa_10 from student_metrics if available, otherwise compute from enrollments
    if "gpa_10" not in df_students.columns or df_students["gpa_10"].isna().all():
        if courses_file.exists() and enrollment_dir.exists():
            df_courses = pd.read_csv(courses_file, usecols=["course_code", "credits"])
            credits_map = dict(zip(df_courses["course_code"], df_courses["credits"]))

            enr_files = sorted(enrollment_dir.glob("enrollment_hk*.csv"))
            if enr_files:
                dfs = [pd.read_csv(f, usecols=["student_id", "course_id", "final_score", "passed"], dtype=str) for f in enr_files]
                df_enr = pd.concat(dfs, ignore_index=True)
                df_enr = df_enr[df_enr["passed"] == "True"]
                df_enr["final_score"] = pd.to_numeric(df_enr["final_score"], errors="coerce")
                df_enr = df_enr[df_enr["final_score"] <= 10.0]
                df_enr["credits"] = df_enr["course_id"].map(credits_map).fillna(0)
                df_enr["weighted_score"] = df_enr["final_score"] * df_enr["credits"]

                grouped = df_enr.groupby("student_id")[["weighted_score", "credits"]].sum()
                grouped["gpa_10"] = grouped["weighted_score"] / grouped["credits"]
                df_students["gpa_10"] = df_students["student_id"].map(grouped["gpa_10"])
            else:
                df_students["gpa_10"] = np.nan
        else:
            df_students["gpa_10"] = np.nan

    return df_students, df_user_st


def build_program_characteristic_1_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    df_students: pd.DataFrame,
    total_univ_students: int,
    program_code: str,
    program_title: str,
    program_desc: str,
) -> None:
    """Build the clean, focused validation report for Characteristic 1: Graduation & Degree Scores for a specific program."""
    ws.views.sheetView[0].showGridLines = True

    # Color Palette & Fonts
    navy_banner = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    navy_header = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
    ice_blue_section = PatternFill(start_color="E8F1F5", end_color="E8F1F5", fill_type="solid")
    guidance_fill = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    font_banner = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    font_section = Font(name="Arial", size=11, bold=True, color="1F4E78")
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Arial", size=9, bold=True)
    font_data = Font(name="Arial", size=9)
    font_guidance = Font(name="Arial", size=9, italic=True, color="1F4E78")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    prog_students = df_students[df_students["program"] == program_code]
    n_prog = len(prog_students)

    # =========================================================================
    # PHẦN 1: TIÊU ĐỀ, BỐI CẢNH & PHẠM VI
    # =========================================================================
    ws.merge_cells("A1:J1")
    c1 = ws["A1"]
    c1.value = f"BÁO CÁO HỌC VỤ TỔNG QUAN DỮ LIỆU MÔ PHỎNG - {program_title.upper()}"
    c1.font = font_banner
    c1.fill = navy_banner
    c1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:J2")
    c2 = ws["A2"]
    c2.value = f"Bối cảnh: Báo cáo tổng quan phân phối dữ liệu học vụ mô phỏng của {program_desc} tại ĐHBK, phục vụ nghiên cứu và huấn luyện mô hình Cảnh báo học vụ sớm (EWS)."
    c2.font = Font(name="Arial", size=9.5, italic=True, color="262626")
    c2.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 24

    ws.merge_cells("A3:J3")
    c3 = ws["A3"]
    pct_total = (n_prog / total_univ_students * 100) if total_univ_students > 0 else 0.0
    c3.value = f"BÁO CÁO HỌC VỤ: TIẾN ĐỘ TỐT NGHIỆP, XẾP LOẠI ĐIỂM & TỶ LỆ BỎ HỌC (K20 - K26) | Quy mô: {n_prog:,} SV ({pct_total:.1f}% toàn trường)"
    c3.font = Font(name="Arial", size=9.5, bold=True, color="1F4E78")
    c3.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 22

    # =========================================================================
    # PHẦN 2: BẢNG 1 - TỶ LỆ TỐT NGHIỆP ĐÚNG HẠN KHÓA 20, 21, 22
    # =========================================================================
    ws.merge_cells("A5:E5")
    s1 = ws["A5"]
    s1.value = "I. BẢNG 1: TỶ LỆ TỐT NGHIỆP ĐÚNG HẠN (KHÓA K20, K21, K22)"
    s1.font = font_section
    s1.fill = ice_blue_section
    s1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[5].height = 26

    table1_headers = [
        "STT", "Khóa học", "Sĩ số tuyển sinh (SV)", "Tốt nghiệp đúng hạn (4.0 năm)",
        "Tỷ lệ đúng hạn (%)"
    ]
    ws.row_dimensions[6].height = 28
    for c_i, h in enumerate(table1_headers, 1):
        cell = ws.cell(row=6, column=c_i, value=h)
        cell.font = font_header
        cell.fill = navy_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    k20_sub = prog_students[prog_students["cohort"] == "K20"]
    k21_sub = prog_students[prog_students["cohort"] == "K21"]
    k22_sub = prog_students[prog_students["cohort"] == "K22"]

    k20_n, k20_ontime = len(k20_sub), int((k20_sub["graduation_timing"] == "on_time").sum())
    k21_n, k21_ontime = len(k21_sub), int((k21_sub["graduation_timing"] == "on_time").sum())
    k22_n, k22_ontime = len(k22_sub), int((k22_sub["graduation_timing"] == "on_time").sum())

    table1_data = [
        (1, "Khóa K20", k20_n, k20_ontime, (k20_ontime / k20_n) if k20_n > 0 else 0.0),
        (2, "Khóa K21", k21_n, k21_ontime, (k21_ontime / k21_n) if k21_n > 0 else 0.0),
        (3, "Khóa K22", k22_n, k22_ontime, (k22_ontime / k22_n) if k22_n > 0 else 0.0),
    ]

    for r_i, (stt, c_name, n, ontime, pct) in enumerate(table1_data, 7):
        ws.row_dimensions[r_i].height = 22
        is_even = (r_i % 2 == 0)

        c_stt = ws.cell(row=r_i, column=1, value=stt)
        c_stt.alignment = Alignment(horizontal="center", vertical="center")

        c_cohort = ws.cell(row=r_i, column=2, value=c_name)
        c_cohort.alignment = Alignment(horizontal="center", vertical="center")
        c_cohort.font = font_bold

        c_n = ws.cell(row=r_i, column=3, value=n)
        c_n.alignment = Alignment(horizontal="right", vertical="center")
        c_n.number_format = "#,##0"

        c_ontime = ws.cell(row=r_i, column=4, value=ontime)
        c_ontime.alignment = Alignment(horizontal="right", vertical="center")
        c_ontime.number_format = "#,##0"

        c_pct = ws.cell(row=r_i, column=5, value=pct)
        c_pct.alignment = Alignment(horizontal="right", vertical="center")
        c_pct.number_format = "0.0%"
        c_pct.font = font_bold

        for col_idx in range(1, 6):
            cell = ws.cell(row=r_i, column=col_idx)
            cell.border = thin_border
            cell.font = font_bold if col_idx in [2, 5] else font_data
            if is_even: cell.fill = zebra_fill

    # Dòng hướng dẫn cách tính Metric Bảng 1
    ws.merge_cells("A10:E10")
    g1 = ws["A10"]
    g1.value = "💡 Cách tính: Tỷ lệ tốt nghiệp đúng hạn = Số sinh viên tốt nghiệp tại kỳ thứ 8 / Tổng sĩ số sinh viên của khóa."
    g1.font = font_guidance
    g1.fill = guidance_fill
    g1.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[10].height = 24

    # =========================================================================
    # PHẦN 3: BẢNG 2 - TỶ LỆ ĐIỂM & XẾP LOẠI CỦA SINH VIÊN ĐÃ TỐT NGHIỆP (K20, K21, K22)
    # =========================================================================
    ws.merge_cells("A12:J12")
    s2 = ws["A12"]
    s2.value = "II. BẢNG 2: TỶ LỆ ĐIỂM & XẾP LOẠI CỦA SINH VIÊN ĐÃ TỐT NGHIỆP (KHÓA K20, K21, K22)"
    s2.font = font_section
    s2.fill = ice_blue_section
    s2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[12].height = 26

    t2_headers = [
        "STT", "Khóa học", "GPA TB (Hệ 10)", "CPA TB (Hệ 4)",
        "Xuất sắc (>= 3.6)", "Giỏi (3.2 - 3.59)", "Khá (2.8 - 3.19)",
        "TB Khá (2.4 - 2.79)", "Trung bình (2.0 - 2.39)", "Dưới TB (< 2.0)"
    ]
    ws.row_dimensions[13].height = 28
    for c_i, h in enumerate(t2_headers, 1):
        cell = ws.cell(row=13, column=c_i, value=h)
        cell.font = font_header
        cell.fill = navy_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Compute graduation score stats for K20, K21, K22
    k20_grad = k20_sub[k20_sub["graduation_ready"] == True]
    k21_grad = k21_sub[k21_sub["graduation_ready"] == True]
    k22_grad = k22_sub[k22_sub["graduation_ready"] == True]

    def get_deg_stats(df_g):
        exc = int((df_g["degree_classification"] == "excellent").sum())
        good = int((df_g["degree_classification"] == "good").sum())
        fair = int((df_g["degree_classification"] == "fair").sum())
        f_avg = int((df_g["degree_classification"] == "fairly_average").sum())
        avg = int((df_g["degree_classification"] == "average").sum())
        below = int((df_g["degree_classification"] == "below_average").sum())
        cpa_10 = float(df_g["gpa_10"].dropna().mean()) if ("gpa_10" in df_g and len(df_g["gpa_10"].dropna())) else 0.0
        cpa_4 = float(df_g["gpa_4"].dropna().mean()) if len(df_g) else 0.0
        tot_g = len(df_g)
        if tot_g == 0:
            pcts = [0.0] * 6
        else:
            pcts = [exc / tot_g, good / tot_g, fair / tot_g, f_avg / tot_g, avg / tot_g, below / tot_g]
        return tot_g, cpa_10, cpa_4, exc, good, fair, f_avg, avg, below, pcts

    k20_stats = get_deg_stats(k20_grad)
    k21_stats = get_deg_stats(k21_grad)
    k22_stats = get_deg_stats(k22_grad)

    t2_rows = [
        (1, "Khóa K20", k20_stats[1], k20_stats[2], *k20_stats[9]),
        (2, "Khóa K21", k21_stats[1], k21_stats[2], *k21_stats[9]),
        (3, "Khóa K22", k22_stats[1], k22_stats[2], *k22_stats[9]),
    ]

    for r_i, (stt, c_name, gpa10, gpa4, p_exc, p_good, p_fair, p_favg, p_avg, p_below) in enumerate(t2_rows, 14):
        ws.row_dimensions[r_i].height = 22
        is_even = (r_i % 2 == 0)

        vals = [stt, c_name, gpa10, gpa4, p_exc, p_good, p_fair, p_favg, p_avg, p_below]
        for c_i, val in enumerate(vals, 1):
            cell = ws.cell(row=r_i, column=c_i, value=val)
            cell.font = font_data
            cell.border = thin_border
            if is_even: cell.fill = zebra_fill

            if c_i == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_i == 2:
                cell.font = font_bold
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_i in [3, 4]:
                cell.font = font_bold
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.number_format = "0.0%"
                cell.alignment = Alignment(horizontal="right", vertical="center")

    # Dòng hướng dẫn cách tính Metric Bảng 2
    ws.merge_cells("A17:J17")
    g2 = ws["A17"]
    g2.value = "💡 Cách tính: GPA 10 = Điểm TB tích lũy hệ 10; CPA 4 = Điểm TB tích lũy thang 4; Tỷ lệ xếp loại = Số SV tốt nghiệp theo từng mức điểm CPA / Tổng số SV tốt nghiệp của khóa."
    g2.font = font_guidance
    g2.fill = guidance_fill
    g2.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[17].height = 24

    # =========================================================================
    # PHẦN 4: BIỂU ĐỒ TRỰC QUAN CỦA BẢNG 2 (ĐẶT NGAY DƯỚI BẢNG 2 - KHUNG A19)
    # =========================================================================
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "percentStacked"
    chart.overlap = 100
    chart.title = "So sánh Cơ cấu Xếp loại Điểm Tốt nghiệp (Khóa K20, K21, K22)"
    chart.width = 22.0
    chart.height = 10.8
    chart.y_axis.title = "Tỷ lệ sinh viên tốt nghiệp (%)"
    chart.x_axis.title = "Khóa tốt nghiệp"

    cats = Reference(ws, min_col=2, min_row=14, max_row=16)  # K20, K21, K22
    data = Reference(ws, min_col=5, max_col=10, min_row=13, max_row=16)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    ws.add_chart(chart, "A19")

    # =========================================================================
    # PHẦN 5: BẢNG 3 - TỶ LỆ BỎ HỌC TỪNG KHÓA (KHÓA K20 ĐẾN K26 - BẮT ĐẦU TỪ DÒNG 44)
    # =========================================================================
    ws.merge_cells("A44:E44")
    s3 = ws["A44"]
    s3.value = "III. BẢNG 3: TỶ LỆ BỎ HỌC TỪNG KHÓA (KHÓA K20 ĐẾN K26)"
    s3.font = font_section
    s3.fill = ice_blue_section
    s3.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[44].height = 26

    table3_headers = [
        "STT", "Khóa học", "Sĩ số tuyển sinh (SV)", "Sinh viên bỏ học (SV)",
        "Tỷ lệ bỏ học (%)"
    ]
    ws.row_dimensions[45].height = 28
    for c_i, h in enumerate(table3_headers, 1):
        cell = ws.cell(row=45, column=c_i, value=h)
        cell.font = font_header
        cell.fill = navy_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    cohorts_all = ["K20", "K21", "K22", "K23", "K24", "K25", "K26"]
    table3_data = []
    for idx, coh in enumerate(cohorts_all, 1):
        c_sub = prog_students[prog_students["cohort"] == coh]
        tot = len(c_sub)
        if "dropout" in c_sub.columns:
            if c_sub["dropout"].dtype == bool:
                drop_cnt = int((c_sub["dropout"] == True).sum())
            else:
                drop_cnt = int((c_sub["dropout"].astype(str).str.lower() == "true").sum())
        else:
            drop_cnt = 0
        pct = (drop_cnt / tot) if tot > 0 else 0.0
        table3_data.append((idx, f"Khóa {coh}", tot, drop_cnt, pct))

    for r_i, (stt, c_name, n, drop_cnt, pct) in enumerate(table3_data, 46):
        ws.row_dimensions[r_i].height = 22
        is_even = (r_i % 2 == 0)

        c_stt = ws.cell(row=r_i, column=1, value=stt)
        c_stt.alignment = Alignment(horizontal="center", vertical="center")

        c_cohort = ws.cell(row=r_i, column=2, value=c_name)
        c_cohort.alignment = Alignment(horizontal="center", vertical="center")
        c_cohort.font = font_bold

        c_n = ws.cell(row=r_i, column=3, value=n)
        c_n.alignment = Alignment(horizontal="right", vertical="center")
        c_n.number_format = "#,##0"

        c_drop = ws.cell(row=r_i, column=4, value=drop_cnt)
        c_drop.alignment = Alignment(horizontal="right", vertical="center")
        c_drop.number_format = "#,##0"

        c_pct = ws.cell(row=r_i, column=5, value=pct)
        c_pct.alignment = Alignment(horizontal="right", vertical="center")
        c_pct.number_format = "0.0%"
        c_pct.font = font_bold

        for col_idx in range(1, 6):
            cell = ws.cell(row=r_i, column=col_idx)
            cell.border = thin_border
            cell.font = font_bold if col_idx in [2, 5] else font_data
            if is_even: cell.fill = zebra_fill

    # Dòng hướng dẫn cách tính Metric Bảng 3
    ws.merge_cells("A53:E53")
    g3 = ws["A53"]
    g3.value = "💡 Cách tính: Tỷ lệ bỏ học = Số sinh viên bỏ học (dropout = True) / Tổng sĩ số sinh viên của khóa."
    g3.font = font_guidance
    g3.fill = guidance_fill
    g3.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[53].height = 24

    # Column Widths Configuration (10 cột chuẩn khít với bảng số liệu và khung biểu đồ)
    col_widths = {
        1: 6,    # A: STT
        2: 16,   # B: Khóa học
        3: 16,   # C: GPA TB Hệ 10
        4: 15,   # D: CPA TB Hệ 4
        5: 16,   # E: Xuất sắc
        6: 16,   # F: Giỏi
        7: 16,   # G: Khá
        8: 16,   # H: TB Khá
        9: 16,   # I: Trung bình
        10: 16,  # J: Dưới TB
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def export_validation_report(output_path: str = None) -> Path:
    """Generate the complete academic validation workbook with focused reports for all 3 programs."""
    output_file = Path(output_path) if output_path else DEFAULT_OUTPUT_FILE

    print(f"[*] Đang tải dữ liệu học vụ từ {METRICS_DIR}...")
    df_students, df_user_st = load_all_datasets()
    total_univ_students = len(df_students)

    programs = [
        {
            "sheet_title": "Đại trà",
            "program_code": "standard",
            "program_title": "CHƯƠNG TRÌNH CHUẨN (ĐẠI TRÀ)",
            "program_desc": "sinh viên Hệ Đại trà",
        },
        {
            "sheet_title": "Chất lượng cao",
            "program_code": "high-quality",
            "program_title": "CHƯƠNG TRÌNH CHẤT LƯỢNG CAO (CLC)",
            "program_desc": "sinh viên Hệ Chất lượng cao",
        },
        {
            "sheet_title": "Tiên tiến",
            "program_code": "advanced",
            "program_title": "CHƯƠNG TRÌNH TIÊN TIẾN",
            "program_desc": "sinh viên Hệ Tiên tiến",
        },
    ]

    print(f"[*] Khởi tạo Workbook và xây dựng Báo cáo Tổng quan cho 3 chương trình...")
    wb = openpyxl.Workbook()

    for idx, prog in enumerate(programs):
        sheet_title = prog["sheet_title"]
        print(f"    -> Đang xử lý sheet: {sheet_title} ({prog['program_code']})...")
        if idx == 0:
            ws = wb.active
            ws.title = sheet_title
        else:
            ws = wb.create_sheet(title=sheet_title)

        build_program_characteristic_1_sheet(
            ws=ws,
            df_students=df_students,
            total_univ_students=total_univ_students,
            program_code=prog["program_code"],
            program_title=prog["program_title"],
            program_desc=prog["program_desc"],
        )

    wb.save(output_file)
    print(f"\n[✓] Đã xuất thành công báo cáo học vụ cho cả 3 chương trình: {output_file.resolve()}\n")
    return output_file


if __name__ == "__main__":
    export_validation_report()
